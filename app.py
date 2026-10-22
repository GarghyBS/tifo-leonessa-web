from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime
import io
from reportlab.lib.pagesizes import A4, A5
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.lib.utils import simpleXML
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO

app = Flask(__name__)
CORS(app)

# Configurazione database (useremo file SQLite per semplicità iniziale)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///iscritti.db'
app.config['SQLALCHEMY_TRACK_NOTIFICATIONS'] = False
db = SQLAlchemy(app)

# Modello di database
class Iscritto(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(100))
    cognome = db.Column(db.String(100))
    data_nascita = db.Column(db.String(10))
    luogo_nascita = db.Column(db.String(50))
    residenza = db.Column(db.String(200))
    codice_fiscale = db.Column(db.String(16))
    numero_tessera = db.Column(db.String(50))
    data_iscrizione = db.Column(db.DateTime, default=datetime.utcnow)
    firmato = db.Column(db.Boolean, default=False)
    data_firma = db.Column(db.DateTime, nullable=True)

with app.app_context():
    db.create_all()

# Template HTML per il form
@app.route('/', methods=['GET'])
def index():
    return '''<!DOCTYPE html>
<html lang="it">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Modulo di Iscrizione - Associazione</title>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body { font-family: Arial, sans-serif; background: #f5f5f5; padding: 20px; }
        .container { max-width: 800px; margin: 0 auto; background: white; padding: 30px; border-radius: 10px; box-shadow: 0 2px 10px rgba(0,0,0,0.1); }
        h1 { text-align: center; color: #333; margin-bottom: 30px; }
        .form-section { margin-bottom: 25px; padding: 20px; background: #f9f9f9; border-radius: 8px; }
        .form-section h2 { color: #555; margin-bottom: 15px; font-size: 18px; border-bottom: 2px solid #4CAF50; padding-bottom: 10px; }
        .form-group { margin-bottom: 15px; }
        label { display: block; margin-bottom: 5px; font-weight: bold; color: #555; }
        input, select { width: 100%; padding: 10px; border: 1px solid #ddd; border-radius: 5px; font-size: 14px; }
        input:focus { border-color: #4CAF50; outline: none; }
        button { width: 100%; padding: 15px; background: #4CAF50; color: white; border: none; border-radius: 5px; font-size: 16px; cursor: pointer; margin-top: 20px; }
        button:hover { background: #45a049; }
        .admin-section { display: none; margin-top: 40px; padding-top: 30px; border-top: 3px solid #333; }
        .admin-section.active { display: block; }
        table { width: 100%; border-collapse: collapse; margin-top: 20px; }
        th, td { padding: 12px; text-align: left; border-bottom: 1px solid #ddd; }
        th { background: #4CAF50; color: white; }
        tr:hover { background: #f5f5f5; }
        .badge { padding: 5px 10px; border-radius: 3px; font-size: 12px; font-weight: bold; }
        .badge.signed { background: #4CAF50; color: white; }
        .badge.not-signed { background: #ff9800; color: white; }
        .confirm-btn { background: #2196F3; width: auto; padding: 8px 15px; margin: 5px; }
        .confirm-btn:hover { background: #1976D2; }
        .download-btn { background: #FF9800; width: auto; padding: 8px 15px; margin: 5px; }
        .download-btn:hover { background: #F57C00; }
        .delete-btn { background: #f44336; width: auto; padding: 8px 15px; margin: 5px; }
        .delete-btn:hover { background: #d32f2f; }
        .stats { display: flex; justify-content: space-around; margin-top: 20px; }
        .stat-box { text-align: center; padding: 15px; background: #e3f2fd; border-radius: 8px; }
        .stat-number { font-size: 24px; font-weight: bold; color: #2196F3; }
    </style>
</head>
<body>
    <div class="container">
        <h1>📝 Modulo di Iscrizione Associazione</h1>
        <form id="iscrizioneForm">
            <div class="form-section">
                <h2>Dati Personali</h2>
                <div class="form-group">
                    <label for="nome">Nome:</label>
                    <input type="text" id="nome" name="nome" required>
                </div>
                <div class="form-group">
                    <label for="cognome">Cognome:</label>
                    <input type="text" id="cognome" name="cognome" required>
                </div>
                <div class="form-group">
                    <label for="data_nascita">Data di Nascita:</label>
                    <input type="date" id="data_nascita" name="data_nascita" required>
                </div>
                <div class="form-group">
                    <label for="luogo_nascita">Luogo di Nascita:</label>
                    <input type="text" id="luogo_nascita" name="luogo_nascita" required>
                </div>
                <div class="form-group">
                    <label for="residenza">Residenza:</label>
                    <input type="text" id="residenza" name="residenza" required>
                </div>
                <div class="form-group">
                    <label for="codice_fiscale">Codice Fiscale:</label>
                    <input type="text" id="codice_fiscale" name="codice_fiscale" maxlength="16" required>
                </div>
                <div class="form-group">
                    <label for="numero_tessera">Numero Tessera Irreducibili:</label>
                    <input type="text" id="numero_tessera" name="numero_tessera" required>
                </div>
            </div>
            <button type="submit">Invia Iscrizione</button>
        </form>
    </div>
    
    <div class="container admin-section">
        <h1>👤 Gestione Iscrizioni (Admin)</h1>
        <div class="stats">
            <div class="stat-box">
                <div class="stat-number"><span id="total">0</span></div>
                <div>Totale</div>
            </div>
            <div class="stat-box">
                <div class="stat-number" style="color: #4CAF50;"><span id="signed">0</span></div>
                <div>Firmati</div>
            </div>
            <div class="stat-box">
                <div class="stat-number" style="color: #ff9800;"><span id="not_signed">0</span></div>
                <div>Non Firmati</div>
            </div>
        </div>
        <button onclick="showAdmin()" style="margin-top: 20px; background: #9C27B0; width: auto; padding: 10px 20px; font-size: 14px;">Accesso Admin</button>
        
        <table id="iscrittiTable">
            <thead>
                <tr>
                    <th>ID</th>
                    <th>Nome</th>
                    <th>Cognome</th>
                    <th>Data Nascita</th>
                    <th>Luogo Nascita</th>
                    <th>Residenza</th>
                    <th>Codice Fiscale</th>
                    <th>Tessera</th>
                    <th>Firmato</th>
                    <th>Azioni</th>
                </tr>
            </thead>
            <tbody></tbody>
        </table>
        <button onclick="downloadExcel()" style="margin-top: 20px; background: #673AB7; color: white; width: auto; padding: 10px 20px; font-size: 14px;">📊 Scarica Excel</button>
    </div>
    
    <div class="container" id="loginSection">
        <h2>🔐 Accedi come Admin</h2>
        <form id="loginForm" style="display: flex; gap: 10px; flex-wrap: wrap;">
            <input type="text" id="username" placeholder="Username" style="width: 200px;">
            <input type="password" id="password" placeholder="Password" style="width: 200px;">
            <button type="submit" style="width: auto; padding: 10px 20px; background: #2196F3;">Login</button>
        </form>
    </div>

    <script>
        // Dati delle registrazioni
        let iscrizioni = []
        let isAdmin = false
        const ADMIN_USERNAME = "admin"
        const ADMIN_PASSWORD = "admin123"
        
        // Carica iscrizioni
        async function loadIscrizioni() {
            try {
                const response = await fetch('/api/iscrizioni')
                iscrizioni = await response.json()
                renderTable()
            } catch (error) {
                console.error('Errore nel caricamento:', error)
            }
        }
        
        // Invia nuova iscrizione
        document.getElementById('iscrizioneForm').addEventListener('submit', async (e) => {
            e.preventDefault()
            
            const data = {
                nome: document.getElementById('nome').value,
                cognome: document.getElementById('cognome').value,
                data_nascita: document.getElementById('data_nascita').value,
                luogo_nascita: document.getElementById('luogo_nascita').value,
                residenza: document.getElementById('residenza').value,
                codice_fiscale: document.getElementById('codice_fiscale').value.toUpperCase(),
                numero_tessera: document.getElementById('numero_tessera').value,
            }
            
            try {
                const response = await fetch('/api/iscrizioni', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify(data)
                })
                
                if (response.ok) {
                    alert('Iscrizione avvenuta con successo! Riceverai a breve le tue documentazioni.')
                    document.getElementById('iscrizioneForm').reset()
                    loadIscrizioni()
                } else {
                    const error = await response.json()
                    alert('Errore: ' + error.message)
                }
            } catch (error) {
                console.error('Errore:', error)
                alert('Si è verificato un errore. Riprova più tardi.')
            }
        })
        
        // Login admin
        document.getElementById('loginForm').addEventListener('submit', (e) => {
            e.preventDefault()
            const username = document.getElementById('username').value
            const password = document.getElementById('password').value
            
            if (username === ADMIN_USERNAME && password === ADMIN_PASSWORD) {
                isAdmin = true
                document.getElementById('loginSection').style.display = 'none'
                document.querySelector('.admin-section').classList.add('active')
                loadIscrizioni()
            } else {
                alert('Credenziali non corrette')
            }
        })
        
        // Mostra/S nascondi admin
        function showAdmin() {
            document.getElementById('loginSection').style.display = 'block'
            document.querySelector('.admin-section').classList.remove('active')
        }
        
        // Renderizza tabella
        function renderTable() {
            const tbody = document.querySelector('#iscrittiTable tbody')
            tbody.innerHTML = ''
            
            document.getElementById('total').textContent = iscrizioni.length
            document.getElementById('signed').textContent = iscrizioni.filter(i => i.firma).length
            document.getElementById('not_signed').textContent = iscrizioni.filter(i => !i.firma).length
            
            iscrizioni.forEach(iscritto => {
                const tr = document.createElement('tr')
                tr.innerHTML = \`
                    <td>\${iscritto.id}</td>
                    <td>\${iscritto.nome}</td>
                    <td>\${iscritto.cognome}</td>
                    <td>\${iscritto.data}</td>
                    <td>\${iscritto.luo}</td>
                    <td>\${iscritto.res}</td>
                    <td>\${iscritto.codice}</td>
                    <td>\${iscritto.tessera}</td>
                    <td><span class="badge \${iscritto.firma ? 'signed' : 'not-signed'}">\${iscritto.firma ? '✔ Firmato' : '✘ Non firmato'}</span></td>
                    <td>
                        <button class="confirm-btn" onclick="confirmIscrizione(\${iscritto.id})">✔</button>
                        <button class="download-btn" onclick="downloadPDF(\${iscritto.id})">📄</button>
                        <button class="delete-btn" onclick="deleteIscrizione(\${iscritto.id})">🗑</button>
                    </td>
                \`
                tbody.appendChild(tr)
            })
        }
        
        // Conferma firma
        async function confirmIscrizione(id) {
            if (!isAdmin) { alert('Accesso admin richiesto'); return }
            
            try {
                const response = await fetch('/api/iscrizioni/' + id + '/conferma', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'}
                })
                
                if (response.ok) {
                    alert('Firma confermata con successo!')
                    loadIscrizioni()
                } else {
                    const error = await response.json()
                    alert('Errore: ' + error.message)
                }
            } catch (error) {
                console.error('Errore:', error)
                alert('Si è verificato un errore.')
            }
        }
        
        // Scarica PDF
        async function downloadPDF(id) {
            try {
                const response = await fetch('/api/pdf/' + id, {
                    responseType: 'blob'
                })
                const blob = await response.blob()
                const url = window.URL.createObjectURL(blob)
                const a = document.createElement('a')
                a.href = url
                a.download = 'ricevuta_' + id + '.pdf'
                document.body.appendChild(a)
                a.click()
                window.URL.revokeObjectURL(url)
                document.body.removeChild(a)
                alert('PDF scaricato con successo!')
            } catch (error) {
                console.error('Errore:', error)
                alert('Errore nel download PDF.')
            }
        }
        
        // Scarica Excel
        async function downloadExcel() {
            if (!isAdmin) { alert('Accesso admin richiesto'); return }
            
            try {
                const response = await fetch('/api/excel')
                const blob = await response.blob()
                const url = window.URL.createObjectURL(blob)
                const a = document.createElement('a')
                a.href = url
                a.download = 'elenco_iscritti_' + new Date().toISOString().split('T')[0] + '.xlsx'
                document.body.appendChild(a)
                a.click()
                window.URL.revokeObjectURL(url)
                document.body.removeChild(a)
                alert('Excel scaricato con successo!')
            } catch (error) {
                console.error('Errore:', error)
                alert('Errore nel download Excel.')
            }
        }
        
        // Elimina iscrizione
        async function deleteIscrizione(id) {
            if (!isAdmin) { alert('Accesso admin richiesto'); return }
            
            if (!confirm('Sei sicuro di voler eliminare questa iscrizione?')) return
            
            try {
                const response = await fetch('/api/iscrizioni/' + id, {
                    method: 'DELETE',
                    headers: {'Content-Type': 'application/json'}
                })
                
                if (response.ok) {
                    alert('Iscrizione eliminata!')
                    loadIscrizioni()
                } else {
                    const error = await response.json()
                    alert('Errore: ' + error.message)
                }
            } catch (error) {
                console.error('Errore:', error)
                alert('Si è verificato un errore.')
            }
        }
        
        // Inizializza
        loadIscrizioni()
    </script>
</body>
</html>'''

# API - Crea nuova iscrizione
@app.route('/api/iscrizioni', methods=['POST'])
def crea_iscrizione():
    data = request.get_json()
    
    try:
        nuova_iscrizione = Iscritto(
            nome=data['nome'],
            cognome=data['cognome'],
            data_nascita=data['data_nascita'],
            luogo_nascita=data['luogo_nascita'],
            residenza=data['residenza'],
            codice_fiscale=data['codice_fiscale'],
            numero_tessera=data['numero_tessera']
        )
        
        db.session.add(nuova_iscrizione)
        db.session.commit()
        
        return jsonify({'success': True, 'id': nuova_iscrizione.id, 'message': 'Iscrizione creata con successo!'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e), 'message': 'Errore nella creazione iscrizione'}), 400

# API - Ottieni tutte le iscrizioni
@app.route('/api/iscrizioni', methods=['GET'])
def get_iscrizioni():
    try:
        iscrizioni = db.session.query(Iscritto).all()
        data = []
        for i in iscrizioni:
            data.append({
                'id': i.id,
                'nome': i.nome,
                'cognome': i.cognome,
                'data': i.data_nascita,
                'luo': i.luogo_nascita,
                'res': i.residenza,
                'codice': i.codice_fiscale,
                'tessera': i.numero_tessera,
                'firma': i.firma,
                'data': i.data_iscrizione.strftime('%Y-%m-%d'),
                'data_firma': i.data_firma.strftime('%Y-%m-%d') if i.data_firma else None
            })
        return jsonify(data)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# API - Conferma firma iscrizione
@app.route('/api/iscrizioni/<int:id>/conferma', methods=['POST'])
def conferma_firma(id):
    try:
        iscrizione = Iscritto.query.get(id)
        if iscrizione:
            iscrizione.firma = True
            iscrizione.data_firma = datetime.utcnow()
            db.session.commit()
            return jsonify({'success': True, 'message': 'Firma confermata con successo!'})
        else:
            return jsonify({'success': False, 'error': 'Iscrizione non trovata'}), 404
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500

# API - Elimina iscrizione
@app.route('/api/iscrizioni/<int:id>', methods=['DELETE'])
def elimina_iscrizione(id):
    try:
        iscrizione = Iscritto.query.get(id)
        if iscrizione:
            db.session.delete(iscrizione)
            db.session.commit()
            return jsonify({'success': True, 'message': 'Iscrizione eliminata con successo!'})
        else:
            return jsonify({'success': False, 'error': 'Iscrizione non trovata'}), 404
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500

# API - Genera PDF (ricevuta A5)
@app.route('/api/pdf/<int:id>', methods=['GET'])
def genera_pdf(id):
    try:
        iscrizione = Iscritto.query.get(id)
        if not iscrizione:
            return jsonify({'error': 'Iscrizione non trovata'}), 404
        
        buf = BytesIO()
        c = canvas.Canvas(buf, pagesize=A5)
        
        # Intestazione
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, 700, "COMUNICAZIONE DI Ricezione")
        c.setFont("Helvetica", 12)
        c.drawString(50, 670, "Associazione")
        c.drawString(50, 650, "Data: %s" % iscrizione.data_nascita)
        
        # Dati utente
        y = 600
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, "Spett.le", fillColor=(0.2, 0.6, 0.8))
        y -= 20
        c.setFont("Helvetica", 10)
        c.drawString(50, y, "Gentile %s %s," % (iscrizione.nome, iscrizione.cognome))
        y -= 40
        c.setFont("Helvetica", 12)
        c.drawString(50, y, "Con la presente si conferma la ricezione della richiesta di iscrizione")
        y -= 40
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, "Numero Tessera Irreducibili: %s" % iscrizione.numero_tessera)
        y -= 60
        
        # Linea per firma
        c.setFont("Helvetica", 10)
        c.drawString(50, y, "_________________")
        y -= 25
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, "Data di ricezione: %s" % iscrizione.data_iscrizione.strftime('%d/%m/%Y'))
        
        # Spazio per la tua firma
        c.drawString(50, y-50, "____________________")
        c.drawString(400, y-50, "Firma Admin")
        
        c.save()
        
        buf.seek(0)
        return send_file(buf, mimetype='application/pdf')
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

# API - Genera PDF A4 con modulo completo e spazio per firma
@app.route('/api/pdf/<int:id>/completo', methods=['GET'])
def genera_pdf_completo(id):
    try:
        iscrizione = Iscritto.query.get(id)
        if not iscrizione:
            return jsonify({'error': 'Iscrizione non trovata'}), 404
        
        buf = BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        
        # Intestazione
        c.setFont("Helvetica-Bold", 20)
        c.drawString(50, 750, "MODULO DI ISCRIZIONE")
        c.setFont("Helvetica", 14)
        c.drawString(50, 710, "Associazione")
        c.drawString(50, 690, "Data di compilazione: %s" % iscrizione.data_iscrizione.strftime('%d/%m/%Y'))
        c.line(50, 660, 550, 660)
        
        # Sezione 1: Delle iscrizioni
        y = 580
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, y, "DETTAGLIO ISCRIZIONE")
        y -= 30
        c.line(50, y, 550, y)
        y -= 40
        
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, "Dati dell'iscritto:")
        y -= 25
        
        c.setFont("Helvetica", 12)
        fields = [
            ("Nome:", iscrizione.nome),
            ("Cognome:", iscrizione.cognome),
            ("Data di nascita:", iscrizione.data_nascita),
            ("Luogo di nascita:", iscrizione.luogo_nascita),
            ("Residenza:", iscrizione.residenza),
            ("Codice Fiscale:", iscrizione.codice_fiscale),
            ("Numero Tessera:", iscrizione.numero_tessera)
        ]
        
        for label, value in fields:
            c.drawString(50, y, label)
            y -= 20
            c.drawString(300, y, value)
            y -= 35
        
        # Sezione 2: Dichiarazione
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, y, "DICHIARAZIONE")
        y -= 30
        c.line(50, y, 550, y)
        y -= 40
        
        c.setFont("Helvetica", 12)
        c.setFont("Helvetica-Oblique", 12)
        y -= 25
        c.drawString(50, y, "Io sottoscritto/a, %s %s,", % (iscrizione.nome, iscrizione.cognome))
        y -= 35
        c.drawString(50, y, "dichiaro di essere a conoscenza e di accettare di iscrivermi")
        y -= 35
        c.drawString(50, y, "all'Associazione, accettando tutti i termini e condizioni")
        y -= 35
        c.drawString(50, y, "previsti dal regolamento associativo.")
        y -= 55
        
        # Spazio per firma
        y -= 70
        c.setFont("Helvetica", 10)
        c.drawString(50, y, "Firma dell'iscritto:")
        c.drawString(50, y-35, "____________________")
        y -= 30
        c.setFont("Helvetica", 10)
        c.drawString(250, y, "Data: _______")
        
        # Spazio per admin
        y -= 50
        c.line(50, y, 500, y)
        y -= 25
        c.drawString(50, y, "Firma dell'Associazione:")
        c.drawString(50, y-35, "____________________")
        y -= 25
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, "Data: _______")
        
        c.save()
        
        buf.seek(0)
        return send_file(buf, mimetype='application/pdf')
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

# API - Genera Excel con riepilogo
@app.route('/api/excel', methods=['GET'])
def genera_excel():
    try:
        iscrizioni = db.session.query(Iscritto).all()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Elenco Iscrizioni"
        
        # Header
        headers = ['ID', 'Nome', 'Cognome', 'Data Nascita', 'Luogo Nascita', 
                  'Residenza', 'Codice Fiscale', 'Numero Tessera', 
                  'Data Iscrizione', 'Firmato', 'Data Firma']
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
            ws.cell(row=1, column=col).font = Font(bold=True)
        
        # Dati
        for row, iscrizione in enumerate(iscrizioni, 2):
            ws.cell(row=row, column=1, value=iscrizione.id)
            ws.cell(row=row, column=2, value=iscrizione.nome)
            ws.cell(row=row, column=3, value=iscrizione.cognome)
            ws.cell(row=row, column=4, value=iscrizione.data_nascita)
            ws.cell(row=row, column=5, value=iscrizione.luogo_nascita)
            ws.cell(row=row, column=6, value=iscrizione.residenza)
            ws.cell(row=row, column=7, value=iscrizione.codice_fiscale)
            ws.cell(row=row, column=8, value=iscrizione.numero_tessera)
            ws.cell(row=row, column=9, value=iscrizione.data_iscrizione.strftime('%d/%m/%Y'))
            ws.cell(row=row, column=10, value="Sì" if iscrizione.firma else "No")
            ws.cell(row=row, column=11, value=iscrizione.data_firma.strftime('%d/%m/%Y') if iscrizione.data_firma else "")
        
        # Stile
        for col in range(1, 12):
            ws.cell(row=1, column=col).fill = PatternFill(start_color='4CAF50', end_color='4CAF50', fill_type='solid')
            ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
        
        ws.font = Font(name='Arial', size=11)
        
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)